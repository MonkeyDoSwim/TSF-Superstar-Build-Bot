#!/usr/bin/env python3
"""Parse the Twilight BATs WWE Champions tier-list workbook into data/index.json.

The workbook mixes TWO tab layouts (see DATA_NOTES.md):
  * Layout A ("7SB new format")  - Tiers cell marks movesets, cards are
    single rows whose Move Sets cell starts "Trainer 26K - <effect>".
  * Layout B ("old format")      - Tier column holds a grade or a role word,
    movesets are flagged by "N Star Build" notes.

Usage:
    python build_index.py [path/to/workbook.xlsx] [-o data/index.json]

With no path argument the live Google Sheet export is downloaded.
Only dependency: openpyxl.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import urllib.request
from pathlib import Path

import openpyxl

SHEET_ID = "1lxSJgCA5KeJeka_yGVxgj_WBKV1sJZ6M9r6ZoRb6oAk"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

# Real data is 15 columns wide; the sheet XML carries formatting out to
# column ~1029, so always cap reads or openpyxl crawls.
MAX_COL = 16

CLASS_TABS = {
    "M STR": ("M", "Strikers"), "M ACR": ("M", "Acrobats"),
    "M TEC": ("M", "Technicians"), "M TRK": ("M", "Tricksters"),
    "M PH": ("M", "Powerhouses"), "M SB": ("M", "Showboats"),
    "F STR": ("F", "Strikers"), "F ACR": ("F", "Acrobats"),
    "F TEC": ("F", "Technicians"), "F TRK": ("F", "Tricksters"),
    "F PH": ("F", "Powerhouses"), "F SB": ("F", "Showboats"),
}
CLASS_ORDER = ["Strikers", "Acrobats", "Technicians",
               "Tricksters", "Powerhouses", "Showboats"]
ROLES = ("Trainer", "Coach", "Promoter")

# "Trainer 26K - effect text" / "Coach 35K: effect text"
ROLE_PREFIX_RE = re.compile(
    r"^(Trainer|Coach|Promoter)\s+(\d+)\s*K\b\s*[:\u2013\u2014-]?\s*", re.IGNORECASE)
DAMAGE_RE = re.compile(r"Damage:\s*([^\n]+)", re.IGNORECASE)
STAR_BUILD_RE = re.compile(r"^\d+\s*Star Build", re.IGNORECASE)

# Normalized header text -> canonical field name. Headers drift between
# tabs (" Coach 1" with a leading space, "Tag Links & Hot Tag", ...).
HEADER_ALIASES = {
    "era & style": "era", "era": "era",
    "move sets": "moves",
    "tiers": "tier", "tier": "tier",
    "trainer 1": "trainer1", "trainer 2": "trainer2",
    "coach 1": "coach1", "coach 2": "coach2", "coach 3": "coach3",
    "skill plates": "skill_plates",
    "ultimate plates": "ultimate_plates",
    "gear & moments": "gear",
    "reversals to avoid": "reversals",
    "notes": "notes",
    "tag links": "tag_links", "tag links & hot tag": "tag_links",
    "entourage ability": "entourage",
    "gameplay videos": "videos",
}

LIST_FIELDS = ("trainers", "coaches", "skill_plates", "ultimate_plates",
               "gear", "tag_links", "reversals", "videos")
FIELD_COLUMNS = {
    "trainers": ("trainer1", "trainer2"),
    "coaches": ("coach1", "coach2", "coach3"),
    "skill_plates": ("skill_plates",),
    "ultimate_plates": ("ultimate_plates",),
    "gear": ("gear",),
    "tag_links": ("tag_links",),
    "reversals": ("reversals",),
    "videos": ("videos",),
}
PROSE_FIELDS = ("notes", "entourage")


def clean(value) -> str:
    """Cell value -> stripped single string ('' for empty)."""
    if value is None:
        return ""
    return str(value).strip()


def clean_name(value) -> str:
    """Persona name: newline-joined in Layout A, space-joined in Layout B."""
    return re.sub(r"\s+", " ", clean(value).replace("\n", " ")).strip()


def normalize_plate(name: str) -> str:
    """Lookup key for a plate: unify apostrophes/quotes, lowercase."""
    name = name.replace("\u2019", "'").replace("\u2018", "'")
    name = name.replace("\u201c", '"').replace("\u201d", '"')
    return re.sub(r"\s+", " ", name).strip().lower()


def column_map(header_row) -> dict:
    """Map canonical field -> column index for one tab's header row."""
    cols = {}
    for idx, raw in enumerate(header_row):
        key = re.sub(r"\s+", " ", clean(raw).lower())
        field = HEADER_ALIASES.get(key)
        if field and field not in cols:
            cols[field] = idx
    return cols


def cell(row, cols, field) -> str:
    idx = cols.get(field)
    if idx is None or idx >= len(row):
        return ""
    return clean(row[idx])


def split_movesets(rows, cols, layout):
    """Split a build's content rows into movesets (lists of row tuples)."""
    if layout == "A":
        # A new moveset starts wherever the Tiers cell is non-empty.
        groups, current = [], []
        for row in rows:
            if cell(row, cols, "tier") and current:
                groups.append(current)
                current = []
            current.append(row)
        if current:
            groups.append(current)
        return groups
    # Layout B: "N Star Build" note flags a new moveset; else chunk by 3.
    groups, current = [], []
    for row in rows:
        if current and STAR_BUILD_RE.match(cell(row, cols, "notes")):
            groups.append(current)
            current = []
        current.append(row)
    if current:
        groups.append(current)
    if len(groups) <= 1 and len(rows) > 3:
        groups = [rows[i:i + 3] for i in range(0, len(rows), 3)]
    return groups


def aggregate(rows, cols, field):
    """Concat + dedupe a list column across a block's rows, in row order."""
    seen, out = set(), []
    for row in rows:
        for f in FIELD_COLUMNS[field]:
            value = cell(row, cols, f)
            for piece in value.split("\n"):
                piece = piece.strip()
                if piece and piece not in seen:
                    seen.add(piece)
                    out.append(piece)
    return out


def aggregate_prose(rows, cols, field):
    """Prose columns (notes/entourage): dedupe whole cell values, join."""
    seen, out = set(), []
    for row in rows:
        value = cell(row, cols, field)
        if value and value not in seen:
            seen.add(value)
            out.append(value)
    return "\n".join(out)


def parse_build(block, cols, layout, gender, klass):
    """Build one build record from a persona's rows."""
    name = clean_name(block[0][0])
    era = cell(block[0], cols, "era")
    groups = split_movesets(block, cols, layout)

    movesets, moveset_tiers = [], []
    for group in groups:
        moves = [cell(r, cols, "moves") for r in group]
        movesets.append([m for m in moves if m])
        moveset_tiers.append(cell(group[0], cols, "tier"))

    if layout == "A":
        tier_text = moveset_tiers[0] if moveset_tiers else ""
        m = DAMAGE_RE.search(tier_text)
        tier = m.group(1).strip() if m else tier_text.split("\n")[0].strip()
        notes = aggregate_prose(block, cols, "notes")
        # The Tiers cells carry the build notes in this layout ("Risk: ...").
        tier_notes = []
        for t in moveset_tiers:
            rest = DAMAGE_RE.sub("", t).strip(" \n")
            if rest and rest not in tier_notes:
                tier_notes.append(rest)
        notes = "\n".join(x for x in tier_notes + ([notes] if notes else []) if x)
    else:
        tier = cell(block[0], cols, "tier")
        notes = aggregate_prose(block, cols, "notes")

    record = {
        "name": name,
        "gender": gender,
        "class": klass,
        "era_style": era,
        "tier": tier,
        "movesets": movesets,
        "moveset_tiers": moveset_tiers,
        "entourage": aggregate_prose(block, cols, "entourage"),
        "notes": notes,
    }
    for field in LIST_FIELDS:
        values = aggregate(block, cols, field)
        if field == "videos":
            values = [v for v in values if "http" in v]
        record[field] = values
    return record


def parse_card(row, cols):
    """Build one card record from its single row."""
    move_text = cell(row, cols, "moves")
    tier_word = cell(row, cols, "tier")
    role, cost, effect = None, None, move_text
    if tier_word in ROLES:
        role = tier_word
    m = ROLE_PREFIX_RE.match(move_text)
    if m:
        role = role or m.group(1).title()
        cost = int(m.group(2))
        effect = re.sub(r"\s+", " ", move_text[m.end():]).strip()
    return {
        "name": clean_name(row[0]),
        "era_style": cell(row, cols, "era"),
        "role": role,
        "cost_k": cost,
        "effect": effect,
    }


def is_card(block, cols) -> bool:
    """A card is a single content row typed by a role word / role prefix."""
    if len(block) != 1:
        return False
    row = block[0]
    if cell(row, cols, "tier") in ROLES:
        return True
    return bool(ROLE_PREFIX_RE.match(cell(row, cols, "moves")))


def parse_class_tab(ws, gender, klass):
    """Parse one class tab -> (builds, cards)."""
    rows = list(ws.iter_rows(min_row=1, max_col=MAX_COL, values_only=True))
    if not rows:
        return [], []
    cols = column_map(rows[0])
    layout = "A" if "reversals" in cols else "B"

    # Group rows into persona blocks: a block starts where column 0 is set.
    blocks, current = [], []
    for row in rows[1:]:
        if clean(row[0]):
            if current:
                blocks.append(current)
            current = [row]
        elif current:
            current.append(row)
    if current:
        blocks.append(current)

    builds, cards = [], []
    for block in blocks:
        # Content rows = rows with anything in a mapped column (drops the
        # blank separator rows that pad some blocks).
        content = [r for r in block
                   if any(clean(r[i]) for i in cols.values() if i < len(r))
                   or clean(r[0])]
        if not content:
            continue
        if is_card(content, cols):
            card = parse_card(content[0], cols)
            card.update({"gender": gender, "class": klass})
            cards.append(card)
        else:
            builds.append(parse_build(content, cols, layout, gender, klass))
    return builds, cards


TIER_SECTION_RE = re.compile(r"^(S\+|S|A|B|C|D)\s+Tier$", re.IGNORECASE)
EXTRA_SECTIONS = {"cheese", "entourage", "trainers", "coaches", "promoters"}


def parse_tier_list(ws):
    """Parse the Tier List tab -> rankings + coming soon + header prose."""
    rows = list(ws.iter_rows(min_row=1, max_col=13, values_only=True))
    result = {"rankings": {}, "coming_soon": [], "notes": []}

    for row in rows[:6]:
        text = clean(row[1]) if len(row) > 1 else ""
        if text:
            result["notes"].append(text)

    # Class of each data column, from the header row (row 8, index 7).
    col_class = {}
    if len(rows) > 7:
        for idx in range(1, 13):
            label = clean(rows[7][idx]).title()
            if label in [c + "s" for c in CLASS_ORDER]:
                gender = "M" if idx <= 6 else "F"
                col_class[idx] = (gender, label)

    for row in rows:
        if clean(row[0]).lower().startswith("coming soon"):
            seen = set()
            for idx in range(1, 13):
                for piece in clean(row[idx]).split(","):
                    piece = piece.strip()
                    if piece and piece not in seen:
                        seen.add(piece)
                        result["coming_soon"].append(piece)
            break

    current_tier, empty_streak = None, 0
    for row in rows[9:]:
        label = clean(row[0])
        names = [(idx, clean(row[idx])) for idx in range(1, 13)
                 if idx < len(row) and clean(row[idx])]
        if not label and not names:
            empty_streak += 1
            if empty_streak > 30:
                break
            continue
        empty_streak = 0
        if label and (TIER_SECTION_RE.match(label)
                      or label.lower() in EXTRA_SECTIONS):
            current_tier = label
            result["rankings"].setdefault(current_tier, [])
        if current_tier is None:
            continue
        for idx, name in names:
            gender, klass = col_class.get(idx, ("", ""))
            result["rankings"][current_tier].append(
                {"name": name, "class": klass, "gender": gender})
    return result


def build_plate_index(builds):
    """plate -> builds using it, split skill vs ultimate."""
    plates = {}
    for b in builds:
        ref = {"name": b["name"], "class": b["class"],
               "gender": b["gender"], "tier": b["tier"]}
        for kind, field in (("skill", "skill_plates"),
                            ("ultimate", "ultimate_plates")):
            for plate in b[field]:
                key = normalize_plate(plate)
                entry = plates.setdefault(
                    key, {"name": plate, "skill": [], "ultimate": []})
                if ref not in entry[kind]:
                    entry[kind].append(ref)
    return plates


def parse_workbook(path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    builds, cards = [], []
    for tab, (gender, klass) in CLASS_TABS.items():
        if tab not in wb.sheetnames:
            continue
        tab_builds, tab_cards = parse_class_tab(wb[tab], gender, klass)
        builds.extend(tab_builds)
        cards.extend(tab_cards)
    tier_list = (parse_tier_list(wb["Tier List"])
                 if "Tier List" in wb.sheetnames
                 else {"rankings": {}, "coming_soon": [], "notes": []})
    wb.close()

    plates = build_plate_index(builds)
    stats = {
        "builds": len(builds),
        "cards": len(cards),
        "cards_by_role": {r: sum(1 for c in cards if c["role"] == r)
                          for r in ROLES},
        "plates": len(plates),
        "tier_list_entries": sum(len(v) for v in tier_list["rankings"].values()),
        "coming_soon": len(tier_list["coming_soon"]),
    }
    return {
        "sheet_id": SHEET_ID,
        "export_url": EXPORT_URL,
        "builds": builds,
        "cards": cards,
        "plates": plates,
        "tier_list": tier_list,
        "stats": stats,
    }


def download_sheet(dest: Path) -> Path:
    req = urllib.request.Request(EXPORT_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp, open(dest, "wb") as fh:
        fh.write(resp.read())
    return dest


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx", nargs="?", help="local workbook (default: download live export)")
    ap.add_argument("-o", "--out", default=str(Path(__file__).parent / "data" / "index.json"))
    args = ap.parse_args(argv)

    if args.xlsx:
        source = Path(args.xlsx)
    else:
        tmp = Path(tempfile.gettempdir()) / "twilight-bats-live.xlsx"
        print(f"Downloading {EXPORT_URL} ...")
        source = download_sheet(tmp)

    print(f"Parsing {source} ...")
    index = parse_workbook(source)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(index, indent=1, ensure_ascii=False), encoding="utf-8")

    s = index["stats"]
    print(f"Wrote {out}")
    print(f"  builds:        {s['builds']}")
    print(f"  cards:         {s['cards']}  (by role: {s['cards_by_role']})")
    print(f"  plates:        {s['plates']}")
    print(f"  tier entries:  {s['tier_list_entries']}")
    print(f"  coming soon:   {s['coming_soon']}")
    per_class = {}
    for b in index["builds"]:
        per_class[b["class"]] = per_class.get(b["class"], 0) + 1
    print(f"  builds/class:  {per_class}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
